#!/usr/bin/env python3
"""Folha-índice para o contabilista — Water Movements (HQ).

Gera um .xlsx com um separador por mês (todas as linhas de hq_invoices vivas),
um separador "Em falta" (linhas sem documento) e um "Resumo" com fórmulas.
Cada linha leva o link do documento na Drive.

Uso (a partir da raiz do repo, com o Supabase CLI ligado ao projecto):
    python3 .claude/skills/gotadagua-platform/references/contabilista_indice.py 2026 saida.xlsx

Depois de gerar, correr o recalc.py do skill xlsx e carregar o ficheiro para a
pasta "Contabilista <ano>" na Drive (substituindo o anterior).
O Manjar Alentejano NÃO entra aqui: é outra empresa e só vive na Drive.
"""
import json, re, subprocess, sys, tempfile, os, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

YEAR = int(sys.argv[1]) if len(sys.argv) > 1 else dt.date.today().year
OUT = sys.argv[2] if len(sys.argv) > 2 else f"Indice Water Movements {YEAR}.xlsx"
MESES = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho",
         "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]

SQL = f"""
select invoice_date, company, coalesce(supplier_nif,'') as nif,
       coalesce(invoice_number,'') as numero, coalesce(description,'') as descricao,
       amount, coalesce(currency,'EUR') as moeda, amount_eur,
       coalesce(category_name,'') as categoria, location_slug as local,
       coalesce(payment_type,'') as pagamento, coalesce(drive_link,'') as link,
       coalesce(file_name,'') as ficheiro, needs_review as rever,
       coalesce(notes,'') as notas
from hq_invoices
where deleted_at is null
  and invoice_date between '{YEAR}-01-01' and '{YEAR}-12-31'
order by invoice_date, company, amount;
"""


def fetch():
    with tempfile.NamedTemporaryFile("w", suffix=".sql", delete=False) as f:
        f.write(SQL)
        path = f.name
    raw = subprocess.run(["supabase", "db", "query", "--linked", "--file", path],
                         capture_output=True, text=True).stdout
    os.unlink(path)
    m = re.search(r"\[\s*\{.*\}\s*\]", raw, re.S)
    return json.loads(m.group(0)) if m else []


ARIAL = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
LINK = Font(name="Arial", size=10, color="0563C1", underline="single")
FILL = PatternFill("solid", fgColor="1F3A5F")
WARN = PatternFill("solid", fgColor="FFF2CC")
COLS = [("Data", 11), ("Fornecedor", 30), ("NIF", 12), ("N.º documento", 20),
        ("Descrição", 46), ("Valor", 11), ("Moeda", 7), ("Valor EUR", 11),
        ("Categoria", 18), ("Local", 13), ("Pagamento", 14), ("Documento", 14),
        ("Estado", 16), ("Notas", 60)]


def estado(r):
    if not r["link"]:
        return "SEM DOCUMENTO"
    return "A REVER" if r["rever"] else "OK"


def sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    for c, (name, width) in enumerate(COLS, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font, cell.fill = HEAD, FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = width
    for i, r in enumerate(rows, 2):
        vals = [dt.date.fromisoformat(r["invoice_date"]), r["company"], r["nif"], r["numero"],
                r["descricao"][:200], float(r["amount"]), r["moeda"], float(r["amount_eur"]),
                r["categoria"], r["local"], r["pagamento"], None, estado(r), r["notas"][:300]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font = ARIAL
        ws.cell(row=i, column=1).number_format = "DD/MM/YYYY"
        ws.cell(row=i, column=6).number_format = "#,##0.00"
        ws.cell(row=i, column=8).number_format = "#,##0.00"
        d = ws.cell(row=i, column=12)
        if r["link"]:
            d.value, d.hyperlink, d.font = "abrir", r["link"], LINK
        else:
            d.value = "—"
        if estado(r) != "OK":
            ws.cell(row=i, column=13).fill = WARN
    n = len(rows) + 1
    t = n + 2
    ws.cell(row=t, column=5, value="Total (EUR)").font = BOLD
    ws.cell(row=t, column=8, value=f"=SUM(H2:H{n})" if rows else 0).font = BOLD
    ws.cell(row=t, column=8).number_format = "#,##0.00"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{max(n, 2)}"
    return n


def main():
    rows = fetch()
    if not rows:
        sys.exit("Sem linhas: confirmar `supabase link` e o ano.")
    wb = Workbook()
    res = wb.active
    res.title = "Resumo"
    por_mes = {m: [] for m in range(1, 13)}
    for r in rows:
        por_mes[int(r["invoice_date"][5:7])].append(r)
    used = [(m, por_mes[m]) for m in range(1, 13) if por_mes[m]]
    ends = {}
    for m, rs in used:
        ends[m] = sheet(wb, MESES[m - 1], rs)
    sheet(wb, "Em falta", [r for r in rows if not r["link"]])
    sheet(wb, "A rever", [r for r in rows if r["link"] and r["rever"]])

    res["A1"] = f"Water Movements Lda (NIF 515059927) — despesas {YEAR}"
    res["A1"].font = Font(name="Arial", size=13, bold=True)
    res["A2"] = ("Gerado a " + dt.date.today().strftime("%d/%m/%Y") +
                 " a partir da app (HQ). Cada separador mensal tem uma linha por movimento; "
                 "a coluna Documento abre o ficheiro na Drive.")
    res["A2"].font = ARIAL
    heads = ["Mês", "Linhas", "Com documento", "Sem documento", "A rever", "Total EUR", "EUR sem documento"]
    for c, h in enumerate(heads, 1):
        cell = res.cell(row=4, column=c, value=h)
        cell.font, cell.fill = HEAD, FILL
        res.column_dimensions[get_column_letter(c)].width = 20 if c > 1 else 14
    row = 5
    for m, rs in used:
        s, n = f"'{MESES[m - 1]}'", ends[m]
        res.cell(row=row, column=1, value=MESES[m - 1]).font = ARIAL
        res.cell(row=row, column=2, value=f"=COUNTA({s}!B2:B{n})").font = ARIAL
        res.cell(row=row, column=3, value=f"=B{row}-D{row}").font = ARIAL
        res.cell(row=row, column=4, value=f'=COUNTIF({s}!M2:M{n},"SEM DOCUMENTO")').font = ARIAL
        res.cell(row=row, column=5, value=f'=COUNTIF({s}!M2:M{n},"A REVER")').font = ARIAL
        res.cell(row=row, column=6, value=f"=SUM({s}!H2:H{n})").font = ARIAL
        res.cell(row=row, column=7, value=f'=SUMIF({s}!M2:M{n},"SEM DOCUMENTO",{s}!H2:H{n})').font = ARIAL
        for c in (6, 7):
            res.cell(row=row, column=c).number_format = "#,##0.00"
        row += 1
    res.cell(row=row, column=1, value="Total").font = BOLD
    for c in range(2, 8):
        L = get_column_letter(c)
        cell = res.cell(row=row, column=c, value=f"=SUM({L}5:{L}{row - 1})")
        cell.font = BOLD
        if c >= 6:
            cell.number_format = "#,##0.00"
    res.cell(row=row + 2, column=1,
             value="Legenda: OK = tem documento; A REVER = tem documento mas há uma dúvida anotada; "
                   "SEM DOCUMENTO = falta a fatura/recibo.").font = ARIAL
    wb.save(OUT)
    print(OUT, len(rows), "linhas")


if __name__ == "__main__":
    main()
